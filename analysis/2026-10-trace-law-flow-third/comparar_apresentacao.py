"""O que mudou entre o deck gerado e o editado à mão — para levar a edição de volta ao `gerar_apresentacao.py`.

Da máquina da base só saem fotos: em vez de fotografar cada slide, rode este script lá e fotografe a saída. Ele lista,
slide a slide, só as formas diferentes, com os valores exatos (posição e tamanho em polegadas, cor em hex, fonte,
tamanho, negrito, texto): `~` forma alterada (só os campos que mudaram), `+` forma nova, `-` forma removida. As figuras
aparecem só como "imagem" com a posição — o conteúdo delas não é lido. Também grava a saída em `<editado>.mudancas.txt`.

Mudança de tema do PowerPoint (Design → Variantes) não aparece: o gerador fixa fonte e cor em cada texto e forma, então
o tema não os alcança — mude a cor ou a fonte direto na forma.

Uso, da raiz:
    uv run python analysis/2026-10-trace-law-flow-third/comparar_apresentacao.py --gerado <gerado.pptx> --editado <editado.pptx>
"""

import argparse
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Emu

TOL = 0.03   # polegadas: abaixo disso, posição e tamanho contam como iguais


def _cor(fmt):
    try:
        return f"#{fmt.rgb}" if fmt.type is not None else None
    except (AttributeError, TypeError, ValueError):
        return None


def _fundo(sh):
    try:
        return _cor(sh.fill.fore_color) if sh.fill.type == 1 else None   # 1 = sólido
    except (AttributeError, TypeError, ValueError):
        return None


def _texto(sh):
    if sh.has_text_frame:
        return sh.text_frame.text
    if sh.has_table:
        return " | ".join(c.text for row in sh.table.rows for c in row.cells)
    return ""


def _fonte(sh):
    """A fonte do primeiro trecho de texto: nome, tamanho, negrito, cor."""
    if not sh.has_text_frame:
        return None
    for p in sh.text_frame.paragraphs:
        for r in p.runs:
            f = r.font
            return (f.name, f.size.pt if f.size else None, bool(f.bold), _cor(f.color))
    return None


def descrever(sh):
    if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
        tipo = "imagem"
    elif sh.has_table:
        tipo = "tabela"
    elif sh.shape_type == MSO_SHAPE_TYPE.TEXT_BOX:
        tipo = "texto"
    else:
        tipo = "forma"
    return {"tipo": tipo, "x": Emu(sh.left).inches, "y": Emu(sh.top).inches, "w": Emu(sh.width).inches,
            "h": Emu(sh.height).inches, "fundo": _fundo(sh), "fonte": _fonte(sh), "texto": _texto(sh)}


def _geo_igual(a, b):
    return all(abs(a[k] - b[k]) <= TOL for k in "xywh")


def _igual(a, b):
    return (a["tipo"] == b["tipo"] and _geo_igual(a, b) and a["fundo"] == b["fundo"] and a["fonte"] == b["fonte"]
            and a["texto"] == b["texto"])


def _curto(t, n=70):
    t = " ".join(t.split())
    return t if len(t) <= n else t[:n - 1] + "…"


def _fmt_fonte(f):
    if f is None:
        return "—"
    nome, tam, negrito, cor = f
    return f"{nome or '?'} {tam if tam is not None else '?'}{' N' if negrito else ''} {cor or ''}".strip()


def _linha(sinal, d):
    geo = f"x={d['x']:.2f} y={d['y']:.2f} w={d['w']:.2f} h={d['h']:.2f}"
    partes = [f"{sinal}{d['tipo']}", geo]
    if d["fundo"]:
        partes.append(f"fundo={d['fundo']}")
    if d["fonte"]:
        partes.append(_fmt_fonte(d["fonte"]))
    if d["texto"]:
        partes.append(f'"{_curto(d["texto"])}"')
    return "  ".join(partes)


def _diferencas(a, b):
    """O que mudou de a (gerado) para b (editado), campo a campo."""
    out = []
    if not _geo_igual(a, b):
        out.append(f"x={b['x']:.2f} y={b['y']:.2f} w={b['w']:.2f} h={b['h']:.2f} "
                   f"(era {a['x']:.2f} {a['y']:.2f} {a['w']:.2f} {a['h']:.2f})")
    if a["fundo"] != b["fundo"]:
        out.append(f"fundo={b['fundo']} (era {a['fundo']})")
    if a["fonte"] != b["fonte"]:
        out.append(f"fonte {_fmt_fonte(b['fonte'])} (era {_fmt_fonte(a['fonte'])})")
    if a["texto"] != b["texto"]:
        out.append(f'texto "{b["texto"]}" (era "{_curto(a["texto"], 50)}")')
    return out


def comparar_slide(gerado, editado):
    g = [descrever(sh) for sh in gerado.shapes] if gerado is not None else []
    e = [descrever(sh) for sh in editado.shapes] if editado is not None else []
    # 1 · iguais saem dos dois lados
    for d in list(e):
        par = next((x for x in g if _igual(x, d)), None)
        if par is not None:
            g.remove(par)
            e.remove(d)
    pares = []
    # 2 · mesmo tipo e mesmo texto (não vazio): mudou posição, cor ou fonte
    for d in list(e):
        par = next((x for x in g if x["tipo"] == d["tipo"] and x["texto"] and x["texto"] == d["texto"]), None)
        if par is not None:
            pares.append((par, d))
            g.remove(par)
            e.remove(d)
    # 3 · mesmo tipo e mesma posição: mudou o texto, a cor ou a fonte
    for d in list(e):
        par = next((x for x in g if x["tipo"] == d["tipo"] and _geo_igual(x, d)), None)
        if par is not None:
            pares.append((par, d))
            g.remove(par)
            e.remove(d)
    linhas = []
    for a, b in pares:
        linhas.append(f"~{b['tipo']}  " + "  ·  ".join(_diferencas(a, b)))
    linhas += [_linha("+", d) for d in e]
    linhas += [_linha("-", d) for d in g]
    return linhas


def main(argv=None):
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--gerado", type=Path, required=True, help="o .pptx que o gerar_apresentacao.py escreveu")
    p.add_argument("--editado", type=Path, required=True, help="a cópia editada à mão")
    args = p.parse_args(argv)
    g, e = Presentation(args.gerado), Presentation(args.editado)
    gs, es = list(g.slides), list(e.slides)
    saida = [f"gerado: {args.gerado.name} ({len(gs)} slides) · editado: {args.editado.name} ({len(es)} slides)"]
    if (g.slide_width, g.slide_height) != (e.slide_width, e.slide_height):
        saida.append(f"tamanho do slide: {Emu(e.slide_width).inches:.2f} × {Emu(e.slide_height).inches:.2f} pol "
                     f"(era {Emu(g.slide_width).inches:.2f} × {Emu(g.slide_height).inches:.2f})")
    sem = 0
    for i in range(max(len(gs), len(es))):
        linhas = comparar_slide(gs[i] if i < len(gs) else None, es[i] if i < len(es) else None)
        if not linhas:
            sem += 1
            continue
        rot = "" if i < len(gs) and i < len(es) else (" (slide novo)" if i >= len(gs) else " (slide removido)")
        saida += ["", f"SLIDE {i + 1:02}{rot}"] + [f"  {l}" for l in linhas]
    saida += ["", f"{sem} slide(s) sem mudança."]
    texto = "\n".join(saida)
    print(texto)
    destino = args.editado.with_suffix(".mudancas.txt")
    destino.write_text(texto + "\n", encoding="utf-8")
    print(f"\n(gravado em {destino})")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
