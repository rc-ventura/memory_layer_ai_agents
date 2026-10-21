"""Gera a apresentação aos stakeholders (o esqueleto é `apresentacao-esqueleto.md`) com as figuras de uma rodada da base 3.

Roda na máquina da base 3, depois do relatório da rodada (`executar.py relatorio --rodada <pasta>`): lê as figuras que
o painel e as taxas gravaram (`<rodada>/painel/`, `<rodada>/taxas/`, a raiz da rodada) e a série mensal
`taxas/taxas_por_mes.csv`, de onde desenha o painel duplo do slide 5 (erros por mês × taxa) e a tabela mensal. Os
números do texto são os dos documentos versionados — cada slide diz de onde e de quando. Uma figura que falta vira um
marcador com o nome do arquivo esperado: o deck sai inteiro em qualquer máquina. Figuras feitas à mão (recortes de trace
anonimizados e varridos, overviews das bases 2 e 3) entram por --figuras, com os nomes do anexo C.

Confere a série mensal contra os totais do texto (30.141 erros, 283.925 passos): se não bater, avisa no terminal, no
slide 5 e em `figuras.md`.

Escreve `<rodada>/apresentacao/` (fora do git, como toda saída de rodada): o .pptx, `figuras.md` (o que entrou, o que
falta, os avisos) e as figuras desenhadas aqui. Sem rodada (fora da máquina da base), escreve em
`apresentacao/` ao lado deste script (ignorada pelo git), só com marcadores e as figuras da base 1.

Uso, da raiz:
    uv run python analysis/2026-10-trace-law-flow-third/gerar_apresentacao.py [--rodada <pasta>] [--figuras <pasta>] [--saida <arquivo.pptx>]
"""

import argparse
import json
import os
import re
from pathlib import Path

import pandas as pd
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
RESULTADOS = HERE / "pipeline" / "resultados"
ASSETS_BASE1 = ROOT / "analysis/2026-09-trace-law-flow/assets"
ERROS_V2, PASSOS_V2 = 30_141, 283_925   # totais da fonte v2 (02-relatorio-achados.md §1) que o texto do deck usa

def caminho_io(path):
    """I/O no Windows com caminho longo (> 260 caracteres, comum no OneDrive da máquina da base); o mesmo de
    pipeline/mineracao_observada.py. Fora do Windows, só o caminho absoluto."""
    path = Path(os.path.abspath(path))
    text = str(path)
    if os.name == "nt" and not text.startswith("\\\\?\\"):
        text = "\\\\?\\UNC\\" + text[2:] if text.startswith("\\\\") else "\\\\?\\" + text
    return Path(text)


NAVY = RGBColor(23, 38, 57)
INK = RGBColor(28, 45, 62)
TEAL = RGBColor(0, 133, 130)
BLUE = RGBColor(51, 111, 178)
GOLD = RGBColor(219, 157, 65)
RED = RGBColor(194, 89, 73)
GRAY = RGBColor(103, 119, 134)
PALE = RGBColor(245, 248, 249)
WHITE = RGBColor(255, 255, 255)
MIST = RGBColor(222, 230, 234)
LITE = RGBColor(231, 244, 242)
DKCARD = RGBColor(35, 54, 73)
MESES = "jan fev mar abr mai jun jul ago set out nov dez".split()

R = None            # a apresentação em montagem
BUSCA = []          # pastas onde procurar as figuras, em ordem
FIGURAS = []        # (slide, arquivo, caminho ou None, de onde vem)
AVISOS = []
CARIMBO = ""        # "rodada <id> · <data>" das figuras da rodada


# ---------------------------------------------------------------- primitivas

def shape(s, x, y, w, h, fill=WHITE, border=None, kind=MSO_SHAPE.ROUNDED_RECTANGLE):
    obj = s.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    obj.fill.solid()
    obj.fill.fore_color.rgb = fill
    if border is None:
        obj.line.fill.background()
    else:
        obj.line.color.rgb = border
    if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        obj.adjustments[0] = 0.06
    return obj


def text(s, value, x, y, w, h, size=17, color=INK, bold=False, align=None, gap=6):
    """Caixa de texto; uma linha por \\n e **negrito** dentro da linha."""
    box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(0.03)
    frame.margin_top = frame.margin_bottom = Inches(0.02)
    for idx, line in enumerate(value.split("\n")):
        p = frame.paragraphs[0] if idx == 0 else frame.add_paragraph()
        for k, pedaco in enumerate(line.split("**")):
            if not pedaco:
                continue
            run = p.add_run()
            run.text = pedaco
            run.font.name = "Aptos"
            run.font.size = Pt(size)
            run.font.bold = bold or k % 2 == 1
            run.font.color.rgb = color
        p.space_after = Pt(gap)
        if align is not None:
            p.alignment = align
    return box


def slide(title, kicker, source, dark=False, notas="", rodada=False):
    s = R.slides.add_slide(R.slide_layouts[6])
    shape(s, 0, 0, 13.333, 7.5, NAVY if dark else WHITE, kind=MSO_SHAPE.RECTANGLE)
    text(s, "MEMÓRIA PARA AGENTES  /  EVIDÊNCIA → DECISÃO", 0.65, 0.25, 9.6, 0.30, 10, MIST if dark else TEAL, True)
    text(s, kicker.upper(), 0.65, 0.62, 10.7, 0.26, 10, GOLD if dark else GRAY, True)
    text(s, title, 0.65, 0.92, 12.0, 0.95, 26 if len(title) <= 62 else 22, WHITE if dark else NAVY, True)
    shape(s, 0.65, 6.98, 12.0, 0.015, MIST, kind=MSO_SHAPE.RECTANGLE)
    if rodada and CARIMBO:
        source += f"  ·  Figuras: {CARIMBO}."
    text(s, source, 0.68, 7.04, 11.4, 0.36, 8, MIST if dark else GRAY, gap=0)
    text(s, f"{len(R.slides):02}", 12.25, 7.06, 0.4, 0.25, 9, MIST if dark else GRAY, align=PP_ALIGN.RIGHT)
    if notas:
        s.notes_slide.notes_text_frame.text = notas
    return s


def card(s, x, y, w, h, label, body, accent=TEAL, body_size=13, fill=PALE, txt=INK, gap=5):
    shape(s, x, y, w, h, fill)
    shape(s, x, y, 0.07, h, accent, kind=MSO_SHAPE.RECTANGLE)
    if h < 1.3:
        text(s, label, x + 0.22, y + 0.08, w - 0.38, 0.28, 11.5, accent, True)
        text(s, body, x + 0.22, y + 0.38, w - 0.38, h - 0.42, min(body_size, 12), txt, gap=gap)
    else:
        text(s, label, x + 0.22, y + 0.14, w - 0.38, 0.40, 14, accent, True)
        text(s, body, x + 0.22, y + 0.58, w - 0.38, h - 0.66, body_size, txt, gap=gap)


def metric(s, x, y, w, number, caption, color=TEAL, h=1.05):
    shape(s, x, y, w, h, PALE)
    text(s, number, x + 0.16, y + 0.08, w - 0.3, 0.5, 26, color, True)
    text(s, caption, x + 0.17, y + 0.62, w - 0.3, h - 0.66, 11, INK, gap=0)


def caixa(s, value, x, y, w, h, fill=LITE, color=INK, size=13, bold=True):
    shape(s, x, y, w, h, fill)
    text(s, value, x + 0.06, y + (h - size / 72 * 1.25) / 2 - 0.03, w - 0.12, h, size, color, bold,
         PP_ALIGN.CENTER, gap=0)


def arrow(s, x, y, w=0.24, color=TEAL):
    shape(s, x, y, w, 0.22, color, kind=MSO_SHAPE.CHEVRON)


def note(s, value, y=6.43, accent=GOLD, size=10, h=0.47):
    shape(s, 0.65, y, 12.0, h, PALE)
    shape(s, 0.65, y, 0.05, h, accent, kind=MSO_SHAPE.RECTANGLE)
    text(s, value, 0.82, y + 0.05, 11.7, h - 0.06, size, gap=0)


def placeholder(s, x, y, w, h, title, detail):
    shape(s, x, y, w, h, PALE, MIST)
    text(s, "FIGURA A INSERIR", x + 0.2, y + 0.12, w - 0.4, 0.24, 9.5, TEAL, True)
    text(s, title, x + 0.2, y + 0.40, w - 0.4, 0.6, 14 if h < 2.5 else 17, NAVY, True, gap=2)
    text(s, detail, x + 0.2, y + (1.0 if h < 2.5 else 1.15), w - 0.4, max(h - 1.2, 0.3), 9.5, GRAY, gap=2)


def achar(nome):
    for pasta in BUSCA:
        p = pasta / nome
        if p.is_file():
            return p
    return None


def figura(s, nome, x, y, w, h, titulo, origem):
    """Insere a figura inteira, centrada na caixa; sem o arquivo, um marcador com o nome esperado."""
    caminho = achar(nome)
    FIGURAS.append((len(R.slides), nome, caminho, origem))
    if caminho is None:
        placeholder(s, x, y, w, h, titulo, f"{nome} — {origem}")
        return
    with Image.open(caminho) as im:
        iw, ih = im.size
    k = min(w / iw, h / ih)
    pw, ph = iw * k, ih * k
    s.shapes.add_picture(str(caminho), Inches(x + (w - pw) / 2), Inches(y + (h - ph) / 2), Inches(pw), Inches(ph))


def tabela(s, linhas, x, y, w, h, larguras, size=10, destaque=(), direita=()):
    n, m = len(linhas), len(linhas[0])
    t = s.shapes.add_table(n, m, Inches(x), Inches(y), Inches(w), Inches(h)).table
    for j, lw in enumerate(larguras):
        t.columns[j].width = Inches(lw)
    for i in range(n):
        t.rows[i].height = Inches(h / n)
    for i, linha in enumerate(linhas):
        for j, v in enumerate(linha):
            c = t.cell(i, j)
            c.fill.solid()
            c.fill.fore_color.rgb = NAVY if i == 0 else (LITE if i in destaque else (PALE if i % 2 else WHITE))
            c.margin_left = c.margin_right = Inches(0.06)
            c.margin_top = c.margin_bottom = Inches(0.02)
            tf = c.text_frame
            tf.word_wrap = True
            tf.paragraphs[0].text = str(v)
            for p in tf.paragraphs:
                p.font.size = Pt(size)
                p.font.name = "Aptos"
                p.font.bold = i == 0 or i in destaque
                p.font.color.rgb = WHITE if i == 0 else INK
                if j in direita:
                    p.alignment = PP_ALIGN.RIGHT
    return t


def n_br(v, casas=0):
    return f"{v:,.{casas}f}".replace(",", "X").replace(".", ",").replace("X", ".")


def rotulo_mes(v):
    m = re.match(r"(\d{4})-?(\d{2})", str(v).strip())
    return f"{MESES[int(m[2]) - 1]}/{m[1][2:]}" if m else str(v)


# ---------------------------------------------------------------- dados da rodada

def serie_mensal(rodada):
    """A série mensal da rodada (taxas/taxas_por_mes.csv), conferida contra os totais do texto."""
    if rodada is None or not (rodada / "taxas" / "taxas_por_mes.csv").is_file():
        return None
    m = pd.read_csv(rodada / "taxas" / "taxas_por_mes.csv").sort_values("mes")
    erros, passos = int(m.erros.sum()), int(m.actionsteps.sum())
    if (erros, passos) != (ERROS_V2, PASSOS_V2):
        AVISOS.append(f"Série mensal da rodada: {n_br(erros)} erros e {n_br(passos)} passos; o texto do deck usa "
                      f"{n_br(ERROS_V2)} e {n_br(PASSOS_V2)} (fonte v2). Conferir a rodada antes de apresentar.")
    return m


def painel_duplo(m, destino):
    """Slide 5: erros absolutos por mês (barras) × taxa de erro por mês (linha) — os dois do mesmo CSV."""
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    rot = [rotulo_mes(v) for v in m.mes]
    x = range(len(m))
    fig, a1 = plt.subplots(figsize=(10, 4.3))
    a1.bar(x, m.erros, color="#336fb2", width=0.62)
    for i, v in enumerate(m.erros):
        a1.text(i, v, n_br(v), ha="center", va="bottom", fontsize=8.5, color="#1c2d3e")
    a1.set_ylabel("erros no mês (barras)", color="#336fb2")
    a1.set_ylim(0, m.erros.max() * 1.18)
    a1.set_xticks(list(x), rot)
    a2 = a1.twinx()
    a2.plot(list(x), 100 * m.taxa, marker="o", color="#db9d41", linewidth=2)
    for i, v in enumerate(m.taxa):
        a2.annotate(f"{n_br(100 * v, 1)}%", (i, 100 * v), xytext=(0, 8), textcoords="offset points", ha="center",
                    fontsize=8.5, color="#8a5a12")
    a2.set_ylabel("taxa de erro por passo (linha)", color="#8a5a12")
    a2.set_ylim(0, 100 * m.taxa.max() * 1.35)
    for a in (a1, a2):
        a.spines["top"].set_visible(False)
    a1.set_title("O erro absoluto cresce; a taxa oscila com o mix de papéis", fontsize=11, loc="left")
    fig.tight_layout()
    fig.savefig(destino, dpi=200)
    plt.close(fig)


# ---------------------------------------------------------------- os slides

def s01_abertura():
    s = slide("Agentes que melhoram com os próprios erros", "Projeto Nº [ANONIMIZADO]  ·  07/2026 → 07/2027",
              "Fontes: docs/work-plan.md; docs/sub-activity-map.md; discussion/scope-and-terminology-decisions.md §§1–2. "
              "Esta apresentação cobre a evidência dos traces (analysis/).", True,
              "Por que não fine-tune? Memória textual é auditável, reversível e atualizável sem pipeline de treino — e é o "
              "que a plataforma consegue versionar hoje.\nQuando vira produto? A timeline responde: piloto em M4 (mai/27).")
    text(s, "Mecanismo de atualização de memória para agentes de IA generativa com aprendizado por reforço aplicado a "
            "fluxos jurídicos", 0.75, 1.90, 11.9, 0.52, 15, GOLD)
    card(s, 0.75, 2.48, 5.83, 1.78, "O PROBLEMA",
         "Cada execução recomeça do zero: a lição de hoje não chega à execução de amanhã — o mesmo erro volta, no "
         "mesmo papel, meses depois.", GOLD, 13.5, DKCARD, WHITE)
    card(s, 6.70, 2.48, 5.95, 1.78, "A PROPOSTA — AUTO-MELHORA CONTROLADA",
         "Sinais de desempenho (revisão humana, feedback da área) ajustam instruções, exemplos e contexto na "
         "**memória textual** do agente — auditável, versionada, reversível, **sem retreinar o modelo**.",
         TEAL, 13.5, DKCARD, WHITE)
    text(s, "PLANO DE TRABALHO  ·  5 MACROATIVIDADES", 0.75, 4.44, 6.5, 0.3, 11, MIST, True)
    plan = [("M1", "Fundamentação +\nPOCs comparativas", "out/26"), ("M2", "Arquitetura +\nvalidação local", "dez/26"),
            ("M3", "Implementação\ndo mecanismo", "mar/27"), ("M4", "Piloto A/B\nem produção", "mai/27"),
            ("M5", "Consolidação +\nexpansão", "jul/27")]
    for i, (m, desc, prazo) in enumerate(plan):
        x = 0.75 + i * 2.40
        shape(s, x, 4.76, 2.26, 1.16, DKCARD)
        text(s, m, x + 0.12, 4.86, 0.5, 0.3, 13, GOLD, True)
        text(s, desc, x + 0.55, 4.84, 1.66, 0.6, 10.5, WHITE, True, gap=0)
        text(s, prazo, x + 0.12, 5.58, 2.0, 0.3, 10.5, MIST)
    shape(s, 0.75, 6.14, 11.9, 0.66, RGBColor(42, 66, 87))
    text(s, "ESTA APRESENTAÇÃO — a evidência empírica", 0.95, 6.26, 4.2, 0.4, 12.5, GOLD, True)
    text(s, "A mineração dos traces reais produz os sinais de erro (Subs 1.4–1.5) que fundamentam a arquitetura (M2) "
            "e o piloto A/B (M3–M4).", 5.15, 6.22, 7.35, 0.55, 11.5, WHITE, gap=0)


def s02_anatomia():
    s = slide("Anatomia do trace: onde mora o erro e por que ele é esquecido", "Como o agente trabalha",
              "Fontes: analysis/2026-09-trace-law-flow/docs/05-schema.md e assets/08-overview-granularidade.png (base 1); "
              "third/docs/01-racionais.md (a janela do erro, base 3).",
              notas="Onde entra a memória? O prompt lê a memória no início da execução — a proposta é atualizar esse "
                    "texto (amarra com o slide 1).\nFechamento: erro no passo, mensagem no próximo prompt, execução "
                    "acaba, tudo descartado. O próximo slide mede exatamente isso.")
    cadeia = [("prompt + memória", GOLD, WHITE), ("TaskStep", LITE, INK), ("PlanningStep", LITE, INK),
              ("ActionStep", BLUE, WHITE), ("ActionStep", BLUE, WHITE), ("…", MIST, INK),
              ("resposta final", TEAL, WHITE)]
    for i, (rot, cor, tc) in enumerate(cadeia):
        x = 0.65 + i * 1.74
        caixa(s, rot, x, 2.0, 1.5, 0.62, cor, tc, 12)
        if i < len(cadeia) - 1:
            arrow(s, x + 1.52, 2.2, 0.18)
    card(s, 0.65, 2.9, 5.8, 1.6, "ZOOM NO ACTIONSTEP",
         "código gerado · pensamento · ferramenta chamada · observação\nÉ o único passo que executa — TaskStep e "
         "PlanningStep não têm campo de erro.", BLUE, 12.5)
    card(s, 6.6, 2.9, 6.05, 1.6, "SE O PASSO ERRA",
         "exceção → a mensagem entra no contexto do próximo passo → o agente **lê o próprio erro** e reage → segue "
         "até a resposta.", RED, 12.5)
    figura(s, "recorte_slide2.png", 0.65, 4.7, 7.4, 1.58, "Recorte real: o erro no passo [n] e a mensagem no [n+1]",
           "--figuras; base 1 via drill_down.py, o mesmo caso do anexo G; anonimizado e varrido")
    shape(s, 8.25, 4.7, 4.4, 1.58, NAVY)
    text(s, "A EXECUÇÃO TERMINA", 8.45, 4.85, 4.0, 0.3, 12, GOLD, True)
    text(s, "todo o contexto é descartado — a próxima execução começa sem a lição", 8.45, 5.2, 4.0, 1.0, 15, WHITE,
         True)
    note(s, "“Ver o erro” = texto no próximo prompt  ·  “entre execuções” = contexto zerado.", h=0.4)


def s03_necessidade():
    s = slide("O agente pode ler o erro e voltar a errar", "A necessidade",
              "Fonte: analysis/2026-09-trace-law-flow/docs/02-relatorio-achados.md, TL;DR, §3.1 e §4 — base 1 (amostra "
              "de 1.000 execuções).",
              notas="E se o agente se recupera depois? Recupera — 100% terminam com resposta (99,3% com conteúdo). O "
                    "problema é custo + reincidência, não crash.\nO próprio LLM não corrige? Corrige dentro da "
                    "execução; entre execuções o contexto é descartado — é a diferença entre recuperação e "
                    "aprendizagem.\n86,3% é robusto a 3 métodos (84,9–86,3%); por mecanismo, a reincidência é 58 "
                    "(13,5%).")
    linhas = [("498", "erros na amostra", 498, NAVY), ("430 · 86,3%", "a mensagem chega ao passo seguinte", 430, BLUE),
              ("51 · 11,9%", "repetem a mesma assinatura logo depois", 51, RED)]
    for i, (num, rot, v, cor) in enumerate(linhas):
        y = 2.0 + i * 1.08
        text(s, num, 0.75, y, 2.6, 0.5, 24, cor, True)
        text(s, rot, 3.4, y + 0.1, 3.9, 0.4, 13.5, INK)
        shape(s, 0.75, y + 0.58, 6.5, 0.22, MIST, kind=MSO_SHAPE.RECTANGLE)
        shape(s, 0.75, y + 0.58, max(0.06, 6.5 * v / 498), 0.22, cor, kind=MSO_SHAPE.RECTANGLE)
    text(s, "Feedback dentro da execução ≠ aprendizagem entre execuções", 0.75, 5.35, 6.6, 0.9, 18, NAVY, True)
    card(s, 7.65, 2.0, 5.0, 1.2, "CUSTO", "Execução com erro custa **≈3× mais tokens**: mediana 146 mil vs 48 mil.",
         GOLD, 12)
    card(s, 7.65, 3.35, 5.0, 1.2, "TEMPO", "O mecanismo “texto longo em literal” reaparece em 145 execuções, em "
         "**10 dos 11 meses**.", TEAL, 12)
    card(s, 7.65, 4.7, 5.0, 1.2, "NÃO É CRASH", "Todas terminam com resposta: o problema é custo e reincidência, não "
         "falha fatal.", BLUE, 12)
    note(s, "11,9% = reincidência imediata dentro do subconjunto cuja mensagem chegou ao contexto — não é taxa geral de "
            "falha.", h=0.4)


def s04_bases():
    s = slide("Três bases, dois regimes: a amostra explica a causa; o sistema mede a escala", "O que medimos",
              "Fontes: third/docs/05-pedido-queries.md e 02-relatorio-achados.md §§1–7; 2026-09-trace-law-flow/docs/"
              "02-relatorio-achados.md; second/docs/02-relatorio-base2.md. Recortes por hor_dat_inicio.", rodada=True,
              notas="A base 3 responde: erro + código + pensamento; o passo anterior (n−1); a reação (n+1, n+2); custo "
                    "da janela; taxa real (10,6% dos passos); concentração por papel e mês; cascatas (execuções com até "
                    "26 erros). Parcial: recuperação (o sinal é sinal, não verdade). Não responde: prompt, modelo, "
                    "trajetória, sucesso final, nada depois de n+2.\nBase 2: “lote de ago/2026” é o nome da extração; o "
                    "recorte é dez/25–ago/26.")
    card(s, 0.65, 1.95, 4.3, 3.0, "AMOSTRAS PROFUNDAS · BASES 1 E 2",
         "**Base 1** — 1.000 execuções · out/25–ago/26\n5.781 passos · 498 erros\n**Base 2** — 1.000 execuções · "
         "dez/25–ago/26\n3.907 passos · 479 erros\nTrace completo: prompt, código, trajetória e resposta — **a causa "
         "do erro está no trace**.", BLUE, 11.5, gap=3)
    shape(s, 5.1, 1.95, 7.55, 3.0, PALE)
    shape(s, 5.1, 1.95, 0.07, 3.0, GOLD, kind=MSO_SHAPE.RECTANGLE)
    text(s, "SISTEMA ATUAL · BASE 3 ★  ·  nov/25–ago/26", 5.32, 2.07, 7.0, 0.3, 14, GOLD, True)
    figura(s, "taxas_funil.png", 5.3, 2.45, 4.1, 2.4, "Funil da rodada",
           "taxas/ da rodada (executar.py taxas): 71.881 execuções → 283.925 passos → 30.141 erros")
    text(s, "CADA LINHA = UM ERRO + JANELA", 9.6, 2.5, 3.0, 0.3, 10.5, GRAY, True)
    for i, rot in enumerate(["n−1", "n", "n+1", "n+2"]):
        caixa(s, rot, 9.6 + i * 0.74, 2.85, 0.66, 0.48, RED if rot == "n" else LITE, WHITE if rot == "n" else INK, 12)
    text(s, "código · pensamento · ferramenta · observação · tokens da janela · sinal de recuperação\n"
            "**30.141 erros** em 20.442 execuções (28,4% das 71.881)", 9.6, 3.45, 3.0, 1.4, 10.5, INK, gap=4)
    tabela(s, [["O que cada fonte vê", "prompt / modelo", "trajetória + resposta", "reação ao erro",
                "população + taxa real"],
               ["Amostras (bases 1 e 2)", "✓", "✓", "✓ completa", "✗ amostra de 1.000"],
               ["Base 3", "✗", "✗", "✓ até n+2", "✓ 10,6% dos passos"]],
           0.65, 5.08, 12.0, 1.2, [2.8, 2.0, 2.4, 2.2, 2.6], 11)
    note(s, "Linha = ActionStep com erro · janela só de ActionSteps · sinal de recuperação ≠ sucesso · as bases não se "
            "somam · só vê erros que levantaram exceção (as silenciosas ficam fora).", h=0.42, size=9.5)


def s05_escala(serie, saida_fig):
    s = slide("Escala: o problema é grande — e concentrado o bastante para ser atacável", "Base 3 · população",
              "Fontes: third/docs/02-relatorio-achados.md §§1, 7.4, 7.5; 04-roadmap.md; série mensal = "
              "taxas/taxas_por_mes.csv da rodada.", rodada=True,
              notas="Por que a taxa caiu de 16,4% para 4,1% e voltou a 15,7%? Os erros absolutos não caíram (1.086 jan → "
                    "1.797 jun → 17.485 ago); o denominador cresceu mais rápido — entraram papéis novos e mais leves "
                    "(junho: ×4,2 execuções de maio). É diluição por mix, não melhora. A prova definitiva é a taxa por "
                    "papel fixo — um dos dados do pedido.\nO try/except explica parte de agosto (sem ele: 10,6% × 6,6% "
                    "jul), não todo.\n15,9% é economia? Não — é a fatia do gasto em passos com erro.\nA amarra: o problema "
                    "não se dilui com o crescimento — cresce junto.")
    metric(s, 0.65, 1.95, 3.9, "30.141", "erros · 10,6% dos 283.925 passos", RED)
    metric(s, 4.7, 1.95, 3.9, "28,4%", "das 71.881 execuções têm erro (20.442)", GOLD)
    metric(s, 8.75, 1.95, 3.9, "592,9 mi", "tokens em passos com erro · 15,9% dos 3.729,3 mi", BLUE)
    if serie is not None:
        painel_duplo(serie, saida_fig / "painel_erros_taxa_mes.png")
    figura(s, "painel_erros_taxa_mes.png", 0.65, 3.15, 7.35, 3.15, "Painel duplo: erros por mês × taxa",
           "desenhado por este script de taxas/taxas_por_mes.csv da rodada")
    if serie is None:
        placeholder(s, 8.15, 3.15, 4.5, 3.15, "Tabela mês × passos × erros × taxa",
                    "taxas/taxas_por_mes.csv da rodada (executar.py taxas)")
    else:
        linhas = [["Mês", "Passos", "Erros", "Taxa"]]
        destaque = []
        for _, r in serie.iterrows():
            rot = rotulo_mes(r.mes)
            if rot[:3] in ("jan", "jun", "ago"):
                destaque.append(len(linhas))
            linhas.append([rot, n_br(r.actionsteps), n_br(r.erros), f"{n_br(100 * r.taxa, 1)}%"])
        tabela(s, linhas, 8.15, 3.15, 4.5, min(3.15, 0.26 * len(linhas)), [1.0, 1.3, 1.1, 1.1], 9.5, destaque,
               (1, 2, 3))
    if AVISOS:
        shape(s, 0.65, 3.15, 7.35, 0.85, RED)
        text(s, "ATENÇÃO — " + " ".join(AVISOS), 0.75, 3.18, 7.15, 0.8, 8.5, WHITE, True, gap=0)
    note(s, "Concentração: agosto/26 = 58% dos erros observáveis (17.031 de 29.392) · 2 papéis ≈ 60% das candidatas.  "
            "A taxa mistura papéis e prompts — não é causal; erro é erro (inclui o sinal de harness).", size=9.5)


def s06_metodo():
    s = slide("O método: dois eixos de leitura — um descreve, outro decide", "Taxonomia e genealogia",
              "Fontes: 2026-09-trace-law-flow/docs/09-metodologia-erro-a-memoria.md; discussion/hipoteses/"
              "trace-error-taxonomy-methodology/; third/docs/01-racionais.md §8.", rodada=True,
              notas="Foi regex? Regras determinísticas sobre o texto e a linha rejeitada, sim — de propósito: cada "
                    "classificação reabre no caso cru e a mesma regra roda igual em qualquer base.\nPor que não um LLM? "
                    "A taxonomia é o instrumento de medida; medir com um modelo seria caro, não reprodutível e mais uma "
                    "coisa a validar.\nPor que dois eixos? A mesma reclamação esconde causas diferentes (1:N) e causas "
                    "diferentes convergem para a mesma lição (N:1).\nDe onde saiu? Do trace; TRAIL, AgentDebug e MAST "
                    "foram bancada de teste e vocabulário.\nCaso real: “resposta sem bloco de código” abriu em seis "
                    "mecanismos com donos diferentes — o sintoma não decide o dono, o mecanismo decide.")
    shape(s, 0.65, 1.95, 7.45, 1.38, PALE)
    text(s, "EIXO SINTOMA — DESCREVE (censo e proveniência)", 0.82, 2.02, 7.0, 0.3, 11.5, GRAY, True)
    for i, rot in enumerate(["30.141 erros", "assinatura", "família"]):
        caixa(s, rot, 0.85 + i * 2.15, 2.38, 1.8, 0.46, MIST, INK, 12)
        if i < 2:
            arrow(s, 2.72 + i * 2.15, 2.5, 0.2, GRAY)
    text(s, "o classify() lê só a mensagem · a família tem um balde que detecta erro novo · **não decide nada**",
         0.82, 2.92, 7.1, 0.36, 10.5, INK, gap=0)
    shape(s, 0.65, 3.43, 7.45, 1.62, LITE)
    text(s, "EIXO MECANISMO — DECIDE (lê a mensagem + a linha de código rejeitada)", 0.82, 3.5, 7.1, 0.3, 11.5, TEAL,
         True)
    passos = [("mecanismo", "1:N — mesma reclamação, causas diferentes"),
              ("unidade", "N:1 — uma frase de memória previne todas?"),
              ("triagem", "é memória? conteúdo único? recorre?"), ("destino", "memória · harness · infra · fila")]
    for i, (rot, sub) in enumerate(passos):
        x = 0.85 + i * 1.8
        caixa(s, rot, x, 3.87, 1.5, 0.44, TEAL, WHITE, 12)
        text(s, sub, x, 4.36, 1.55, 0.66, 9, INK, gap=0)
        if i < 3:
            arrow(s, x + 1.53, 3.98, 0.18)
    figura(s, "genealogia_observada.png", 8.25, 1.95, 4.4, 3.1, "Genealogia da rodada (Sankey)",
           "painel/ da rodada (executar.py painel); versão da base 1 no anexo A")
    card(s, 0.65, 5.15, 12.0, 1.15, "EXEMPLO · “COULD NOT INDEX”",
         "O classify() vê **1 sintoma**; o submecanismo() separa **3 causas** lendo a linha rejeitada — dict tratado "
         "como lista · string tratada como dict · chave inexistente: 3 consertos diferentes (caso real no anexo H).",
         GOLD, 12)
    note(s, "Prova: cobertura 97,5% (a regra decide em 29.392 dos 30.141) · resíduo 1,9% · auditoria independente com 0 "
            "divergências · candidata recorre em ≥3 execuções e ≥2 meses · sem LLM: auditável e reproduzível.",
         size=9.5, h=0.42)


def s07_destinos():
    s = slide("Para onde vão os erros: lição do agente × ajuste nosso", "Destinos",
              "Fontes: third/docs/02-relatorio-achados.md §3, §7.2, §7.8–7.10; 04-roadmap.md; dossiê do sandbox §2h. "
              "O callout é a foto de 14/10.", rodada=True,
              notas="Por que sinal de harness não vira memória? Porque a instrução está errada: escrever memória seria "
                    "ensinar o agente a desobedecer o próprio prompt. O conserto é na esteira, e a squad pode fazer "
                    "já.\nDiferença para harness/infra: no sinal de harness o erro é do agente (obedeceu uma instrução "
                    "ruim); em harness/infra falhou o protocolo ou a plataforma.\nApoio: 3 candidatas concentram 75% dos "
                    "erros das candidatas e 85% dos tokens delas (14/10).\n“Resposta sem bloco de código”: família "
                    "Protocolo do harness → unidade de 4.592 erros, a maior não-memória; seis mecanismos com donos "
                    "diferentes.\nCandidata já vira memória? Não — degrau 1 de 4 (slide 8).")
    figura(s, "destinos.png", 0.65, 1.95, 7.25, 3.35, "Destinos em erros × tokens",
           "painel/ da rodada (executar.py painel)")
    tabua = [("MEMÓRIA", TEAL, "lição que o agente pode aprender → texto na memória. Dono: o projeto (degrau 1 de 4)."),
             ("SINAL DE HARNESS", GOLD, "erro do agente induzido pela nossa instrução — conserto no prompt ou no "
                                        "contrato (squad)."),
             ("HARNESS / INFRA", GRAY, "o modelo viola o formato que o prompt ensina, ou a plataforma falha (sem bloco "
                                       "de código, timeout, LLM upstream) — squad."),
             ("RESÍDUO · PENDENTE · CRÍTICO", RED, "a regra não reconheceu, faltou dado ou a execução morreu — filas "
                                                   "da taxonomia e da query.")]
    for i, (rot, cor, corpo) in enumerate(tabua):
        y = 1.95 + i * 0.85
        shape(s, 8.05, y, 4.6, 0.78, PALE)
        shape(s, 8.05, y, 0.06, 0.78, cor, kind=MSO_SHAPE.RECTANGLE)
        text(s, rot, 8.22, y + 0.04, 4.3, 0.25, 10, cor, True, gap=0)
        text(s, corpo, 8.22, y + 0.28, 4.35, 0.5, 9, INK, gap=0)
    card(s, 0.65, 5.42, 12.0, 0.95, "O MAIOR ERRO ISOLADO É NOSSO",
         "5.764 erros e 112 mi tokens (foto de 14/10): o prompt do ContestacaoCivel manda verificar a variável com "
         "try/except, que o sandbox não suporta — o agente obedece e quebra. **Um ajuste de prompt elimina ~19% dos "
         "erros antes de qualquer piloto de memória.**", GOLD, 11)
    note(s, "Depois de 14/10, a mesma causa ganhou mais 2.351 erros vindos do sandbox (Ajustes 21 e 27; slide 9) — a "
            "figura da rodada já os mostra no sinal de harness.  Candidata = degrau 1 de 4 (slide 8).", size=9.5,
         h=0.42, y=6.47)


def s08_processo():
    s = slide("Como uma candidata vira memória: o processo e a escada", "Governança",
              "Fontes: third/docs/06-procedimento-mineracao-candidatas.md; 2026-09-trace-law-flow/docs/"
              "17-validacao-memorias.md; .claude/skills/mineracao-base/referencias/kit.md; pocs/blueprint-m1-pocs-minimos.md.",
              notas="Por que inventário antes de destino? Não dá para rotear o que não foi contado; o inventário revela "
                    "as partes com donos diferentes (slide 9).\nA auditoria audita o quê? Um script independente "
                    "recalcula cada medida; o encontro compara medida a medida.\nPor que 4 degraus? Cada um responde "
                    "uma pergunta com um instrumento: existência ≠ coerência ≠ causa ≠ eficácia.\nQuando injeta de "
                    "verdade? Depois do degrau 3, só na POC mínima controlada.\nE se falhar num degrau? Volta para o "
                    "dono certo: falha de ambiente ou conflito vira ajuste de prompt; falha de recorrência vai para a "
                    "régua de memória rara.")
    text(s, "O PIPELINE — DE UNIDADE A CANDIDATA  (detalhe: anexo I)", 0.65, 1.92, 9.0, 0.3, 11.5, GRAY, True)
    etapas = [("TRIAGEM", "a taxonomia roteia os erros nas unidades"),
              ("INVENTÁRIO", "um script conta o que os erros têm em comum → as partes com dono próprio"),
              ("INVESTIGAÇÃO", "a hipótese é testada nas bases; a evidência decide o destino de cada parte"),
              ("ARQUIVO FINAL", "o script monta a orientação e o catálogo → CANDIDATA catalogada")]
    for i, (rot, corpo) in enumerate(etapas):
        x = 0.65 + i * 3.08
        shape(s, x, 2.27, 2.8, 1.3, LITE)
        text(s, rot, x + 0.15, 2.35, 2.5, 0.3, 13, TEAL, True)
        text(s, corpo, x + 0.15, 2.7, 2.55, 0.85, 10.5, INK, gap=0)
        if i < 3:
            arrow(s, x + 2.83, 2.82, 0.2)
    text(s, "A auditoria independente mede junto em cada etapa; o encontro confere medida a medida (sandbox, 21/10: "
            "**77 de 77 iguais**).", 0.65, 3.65, 12.0, 0.32, 11, INK, gap=0)
    text(s, "A ESCADA — QUANDO A MEMÓRIA ESTÁ PRONTA PARA INJETAR  (detalhe: anexo J)", 0.65, 4.1, 10.0, 0.3, 11.5,
         GRAY, True)
    degraus = [("1 · CANDIDATA", "O erro existe assim e com essa frequência?", "determinístico (a mineração)", BLUE),
               ("2 · CONFERIDA", "A lição é coerente com o ambiente e o prompt?",
                "5 checagens por script: ambiente, conflito, repete, recorrência, uma causa", GOLD),
               ("3 · REVISADA", "A causa é a que a regra supõe?",
                "2 leitores às cegas + gate do pesquisador → PRONTA PARA INJETAR", TEAL),
               ("4 · VALIDADA", "Injetada, reduz o erro sem criar outro?",
                "experimento antes × depois — a POC mínima", RED)]
    for i, (rot, perg, inst, cor) in enumerate(degraus):
        x = 0.65 + i * 3.08
        shape(s, x, 4.45, 2.8, 1.85, PALE)
        shape(s, x, 4.45, 2.8, 0.08, cor, kind=MSO_SHAPE.RECTANGLE)
        text(s, rot, x + 0.15, 4.6, 2.5, 0.3, 13, cor, True)
        text(s, perg, x + 0.15, 4.95, 2.55, 0.65, 11.5, INK, True, gap=0)
        text(s, inst, x + 0.15, 5.6, 2.55, 0.65, 9.5, GRAY, gap=0)
        if i < 3:
            arrow(s, x + 2.83, 5.25, 0.2, cor)
    note(s, "Até o degrau 4, a memória é hipótese com evidência: “catalogada” prova que o erro existe — não que a "
            "memória ensina.", h=0.4, y=6.47)


def s09_sandbox():
    s = slide("Caso sandbox: um rótulo, várias causas — e a maior era o nosso prompt", "Caso real · em fechamento",
              "Fontes: third/propostas/U_sandbox/dossie_base3.md §§2, 2b, 2d, 2h (itens 8–22); analysis/decisoes.md "
              "19–21/10; analysis/registros/memorias.json (rodada 99e032ad, 21/10).",
              notas="Por que cada grupo foi para aquele destino? Se a instrução veio do prompt, o conserto é no harness; "
                    "se é do agente em qualquer papel, memória global; se depende da lista daquele papel, memória do "
                    "papel.\nA verificação: em julho o agente verificava com globals(); em agosto, com try/except — o "
                    "prompt passou a ensinar o jeito e o erro trocou de sintoma, não sumiu (anexo L).\nO "
                    "get_document_text: o prompt do OBFCivel manda usá-lo, mas ele não está entre as ferramentas "
                    "declaradas — em duas bases e todas as versões.\nAs 17 já são memória injetada? Não: catalogadas, "
                    "degrau 1. Três passaram o degrau 2; as decisões abertas (B do OBFCivel, o managerAgent na C "
                    "global, a parte F) seguem com a unidade.\nO json no RoteadorCivel: o import funciona no papel e o "
                    "agente se corrige no passo seguinte em 9 de 10 — é esquecimento do agente, memória.")
    shape(s, 0.65, 1.95, 12.0, 0.55, LITE)
    text(s, "**U_sandbox** — “o agente escreve código com um nome que o sandbox recusa”  ·  **3.653 erros** em 14/10  ·  "
            "a 2ª maior candidata da triagem", 0.82, 2.06, 11.7, 0.35, 12.5, INK, gap=0)
    card(s, 0.65, 2.62, 4.75, 2.6, "ERRO NOSSO → HARNESS · 2.518 (69%)",
         "**Verificação de existência** — ContestacaoCivel: 2.217 + 134\no prompt manda “verifique se a variável "
         "existe”; o agente tenta globals() e try, e o sandbox barra.\n**Ferramenta citada sem ser declarada** — "
         "OBFCivel: 167\no prompt manda usar get_document_text, que não está entre as ferramentas.\n+ 35 eval por "
         "outra regra.", RED, 10.5, gap=3)
    text(s, "FICA NO SANDBOX · 1.100 erros · 17 memórias (rodada de 21/10)", 5.6, 2.62, 7.0, 0.3, 11.5, TEAL, True)
    partes = [("C · json, re, datetime sem import", 423, "memória global", TEAL),
              ("E · depuração proibida (repr, dir…)", 317, "memória global", TEAL),
              ("C · pd sem import (JoogleAnalytics)", 231, "memória do papel", BLUE),
              ("B · nomes do framework (OBFCivel)", 44, "memória do papel ⚑", BLUE),
              ("D · import fora da lista (10 papéis)", 40, "reforço · monitorada", GRAY),
              ("F · como verificar no sandbox", 37, "em aberto", GOLD),
              ("raras · 2 memórias", 3, "monitoradas", GRAY),
              ("sem nome", 5, "—", MIST)]
    for i, (rot, v, dest, cor) in enumerate(partes):
        y = 2.98 + i * 0.28
        text(s, rot, 5.6, y, 2.85, 0.27, 9.5, INK, gap=0)
        shape(s, 8.5, y + 0.04, max(0.04, 2.2 * v / 423), 0.18, cor, kind=MSO_SHAPE.RECTANGLE)
        text(s, f"**{n_br(v)}**  {dest}", 10.78, y, 1.9, 0.27, 9.5, INK, gap=0)
    card(s, 0.65, 5.33, 5.9, 1.05, "CONCENTRADO NÃO DECIDE O DONO",
         "A verificação (só no ContestacaoCivel) e o pd (231 de 233 no JoogleAnalytics) estão num papel só — e tiveram "
         "destinos opostos. O dono se lê no mecanismo: onde mora a instrução.", GOLD, 10)
    card(s, 6.7, 5.33, 5.95, 1.05, "DEGRAU 2 — A ESCADA APLICADA",
         "3 memórias passam e seguem para a leitura de casos: C global, C do JoogleAnalytics e E — **971 erros, 88% do "
         "que ficou**. As 10 D repetem o prompt; ⚑ B do OBFCivel em decisão.", BLUE, 10)
    note(s, "Em fechamento: rodada 99e032ad (21/10) — 9 passos fecham, cobertura 99,6%, 77 medidas iguais à auditoria.  "
            "Abertos: B do OBFCivel, o managerAgent na C global, a parte F.  Catalogada ≠ injetada.", size=9.5,
         h=0.42, y=6.47)


def s10_custo():
    s = slide("O custo de não aprender: o erro mora na cauda cara", "Quanto vale",
              "Fontes: third/docs/01-racionais.md §7; 02-relatorio-achados.md §7.2, §7.5; 04-roadmap.md; figura de "
              "concentração = taxas/ da rodada. Teto e divisão 387/112 = foto de 14/10.", rodada=True,
              notas="13,4% é pouco? É o que já é atacável pelo loop hoje. O valor maior é que um sistema que não aprende "
                    "erra na proporção do crescimento — e o erro senta na cauda que carrega o gasto.\n50% com erro "
                    "prova causa? Não sozinho: é associação (o papel pesado erra mais). Mas três medidas apontam o "
                    "mesmo lado: 1,6× por passo, ~4,3× por execução, erro ~2× mais presente na cauda cara.\nPor que não "
                    "em R$? Falta o preço por modelo; com ele, a conversão é uma multiplicação.\nE se a eficácia for "
                    "baixa? Com 25%, o orçamento cai a ~341 tokens/passo; com 10%, ~136 — saber isso antes é o "
                    "ponto.\nCONFERIR antes de apresentar: os 28/50/81 do texto contra a figura da rodada.")
    card(s, 0.65, 1.95, 7.25, 1.52, "O ERRO ESCALA E CONCENTRA",
         "**No tempo:** 1.086 erros em jan → 17.485 em ago (16×) — sem memória, o erro cresce com o sistema.\n"
         "**No custo:** execuções com erro — todas 28,4% · 10% mais caras 50% · 1% mais cara 81% (os 10% = 64% do "
         "gasto).", TEAL, 11, gap=3)
    card(s, 0.65, 3.57, 7.25, 1.6, "O QUE É ATACÁVEL HOJE — EM TOKENS, NÃO EM R$",
         "592,9 mi em erros → 387 mi memória + 112 mi sinal de harness = **499 mi = 13,4% do gasto**, dentro do mesmo "
         "loop.\n**A régua:** injetar compensa se custo da memória < eficácia × orçamento. Orçamento máximo: 1.363 "
         "tokens/passo (todas as candidatas) · 729 (só a maior unidade) · com eficácia de 25%, ~341.", GOLD, 11,
         gap=3)
    card(s, 0.65, 5.27, 3.55, 1.05, "ERRO NOSSO → AJUSTE DE PROMPT", "barato — a squad já pode fazer", GOLD, 11.5)
    card(s, 4.35, 5.27, 3.55, 1.05, "ERRO DO AGENTE → MEMÓRIA", "o piloto mede se compensa", TEAL, 11.5)
    figura(s, "taxas_concentracao.png", 8.05, 1.95, 4.6, 4.35, "Curva de concentração do custo",
           "taxas/ da rodada (executar.py taxas): o gradiente 28 → 50 → 81")
    note(s, "Teto = hipótese: conta o passo que falhou, não a cascata de recuperação nem a qualidade da resposta · "
            "execução com erro ~4,3× os tokens (parte é mix de papéis) · a reclassificação do sandbox muda a divisão "
            "387/112; o total de 13,4% se mantém.", size=9, h=0.47)


def s11_conclusoes():
    s = slide("O que os números mandam: quatro conclusões", "Síntese",
              "Fontes: os slides 5, 7, 9 e 10 — nenhum número novo, a leitura dos números já mostrados.",
              notas="E se o erro não causar o custo? Mesmo sem causalidade, a conclusão operacional vale: o conserto "
                    "mirado atinge a cauda cara.\nPor que um slide só de conclusões? Quem concorda com as quatro já "
                    "concordou com o pedido.")
    linhas = [("FATO", NAVY, "O erro não é periferia: é núcleo do custo",
               "10,6% dos passos e 15,9% dos tokens; quanto mais cara a execução, mais certo o erro — 28,4% em todas, 50% "
               "nas 10% mais caras, 81% na 1% mais cara."),
              ("PADRÃO", GOLD, "Uma parte grande é erro nosso, com conserto conhecido",
               "o maior erro isolado e 69% do rótulo sandbox vieram de instruções do prompt; e o prompt anuncia um "
               "ambiente que não é o real (anexo M)."),
              ("IMPLICAÇÃO", TEAL, "O que é lição do agente se concentra — e concentração é alavanca",
               "3 candidatas = 75% dos erros das candidatas e 85% dos tokens delas; 2 papéis ≈ 60%: pouco alvo, muito "
               "retorno potencial."),
              ("AÇÃO", RED, "O piloto mira memórias e papéis, não a média",
               "começar pelas 3 memórias que passaram o degrau 2 (971 erros) e pelos papéis que concentram as "
               "candidatas; corrigir o prompt em paralelo.")]
    for i, (tag, cor, titulo, corpo) in enumerate(linhas):
        y = 1.95 + i * 1.0
        caixa(s, tag, 0.65, y, 1.75, 0.9, cor, WHITE, 13)
        shape(s, 2.5, y, 10.15, 0.9, PALE)
        text(s, titulo, 2.68, y + 0.07, 9.8, 0.33, 14.5, NAVY, True, gap=0)
        text(s, corpo, 2.68, y + 0.43, 9.8, 0.45, 10.5, INK, gap=0)
    text(s, "⇒ o pedido: corrigir já o que é nosso, e testar a memória onde o custo concentra (slide 12)", 0.65, 5.98,
         12.0, 0.4, 15, NAVY, True, gap=0)
    note(s, "Parte do gradiente é mecânica (execução mais longa = mais passos = mais chances de errar): descrição "
            "consistente com 1,6× por passo e ~4,3× por execução, não prova causal.", h=0.4, y=6.47)


def s12_pedido():
    s = slide("O pedido: corrigir o que é nosso, validar a memória antes de implantar", "Decisão",
              "Fontes: 2026-09-trace-law-flow/docs/17-validacao-memorias.md; third/docs/04-roadmap.md; "
              "05-pedido-queries.md; analysis/roadmap-transversal.md; pocs/blueprint-m1-pocs-minimos.md.",
              notas="Por que os avisos vêm antes do piloto? São erros nossos com conserto conhecido: corrigir não "
                    "depende de memória e já reduz o erro.\nO piloto é controlado: as memórias entram só depois do "
                    "degrau 3, numa POC mínima, com o mesmo recorte com e sem memória.\nO que não estamos pedindo: "
                    "implantar memória em produção — isso é M4, depois do degrau 4.")
    blocos = [("01 · PARA A SQUAD DO HARNESS — JÁ", RED,
               "**Corrigir o prompt do ContestacaoCivel:** a instrução de verificar a variável, que o sandbox barra.\n"
               "**Declarar ou renomear o get_document_text** no OBFCivel.\n**Explicar a lista de imports:** a do prompt "
               "não é a do interpretador (531 de 531 imports de json rodam com o prompt dizendo que não)."),
              ("02 · PARA O PROJETO — O PILOTO", TEAL,
               "**Leitura de casos (degrau 3)** das 3 memórias que passaram o degrau 2.\n**POC mínima (degrau 4):** "
               "agente com × sem memória no mesmo recorte — reincidência, tokens, latência e qualidade da resposta."),
              ("03 · OS DADOS QUE FALTAM", GOLD,
               "**Prompt e modelo por papel** na base 3.\n**Passos sem erro** (a janela completa): recuperação e "
               "qualidade.\n**Falhas silenciosas** (o erro que não levanta exceção).\n**Preço por modelo**, para "
               "converter tokens em R$.")]
    for i, (rot, cor, corpo) in enumerate(blocos):
        card(s, 0.65 + i * 4.05, 1.95, 3.9, 3.45, rot, corpo, cor, 11.5, gap=6)
    shape(s, 0.65, 5.55, 12.0, 0.85, NAVY)
    text(s, "DECISÃO SOLICITADA", 0.9, 5.68, 3.0, 0.3, 12, GOLD, True)
    text(s, "Autorizar o piloto controlado e encaminhar os três avisos ao time do harness — a evidência justifica "
            "testar, ainda não justifica prometer.", 3.6, 5.66, 8.9, 0.7, 13.5, WHITE, True, gap=0)


def s13_referencias():
    s = slide("Referências e onde cada número vive", "Fontes",
              "Fichamentos em analysis/2026-09-trace-law-flow/literature/ · papers/reading-queue.md. Os anexos A–M "
              "carregam a evidência específica de cada slide.")
    card(s, 0.65, 1.95, 6.0, 4.4, "BANCADA — A LITERATURA QUE TESTAMOS (slide 6)",
         "**MAST** — Cemri et al., UC Berkeley, NeurIPS’25 · arXiv:2503.13657\n14 modos de falha; nenhum cobre “o "
         "código levantou exceção” — a lacuna que a nossa família mais frequente ocupa.\n**TRAIL** — Patronus AI · "
         "arXiv:2505.08638\ntaxonomia + dataset anotado; deu os detectores de falha silenciosa e o teste de "
         "cobertura.\n**AgentDebug** · arXiv:2509.25370\nroteador módulo → tipo de memória; o mais próximo da tese — "
         "referência, não precedente de memória persistente.\n**ToolScan / SpecTool** — Salesforce, ICLR’25 WS · "
         "arXiv:2411.13547 (+ ToolFailBench, ICML’26 WS)\nos erros de ferramenta e os 4 modos de falha silenciosa.",
         BLUE, 10, gap=3)
    card(s, 6.8, 1.95, 5.85, 4.4, "DOCUMENTOS DO PROJETO — ONDE CADA NÚMERO VIVE",
         "**02-relatorio-achados.md** — os achados da base 3 (slides 4, 5, 7, 10)\n**09-metodologia-erro-a-memoria.md** "
         "— a genealogia (slide 6)\n**17-validacao-memorias.md** — a escada de 4 degraus (slide 8)\n"
         "**dossie_base3.md** — o caso sandbox, achado por achado (slide 9, anexos L e M)\n"
         "**registros/memorias.json** — o catálogo das 16 memórias (slide 9, anexo E)\n**decisoes.md** — as decisões "
         "já tomadas, com data\n**01-racionais.md §7** — o teto e a régua (slide 10)", TEAL, 10, gap=5)


# ---------------------------------------------------------------- anexos

def anexo(letra, titulo, fonte, rodada=False):
    return slide(titulo, f"Anexo {letra}", fonte, rodada=rodada)


def anexos():
    s = anexo("A", "Genealogia completa: o exemplo visual da base 1",
              "Fonte: analysis/2026-09-trace-law-flow/assets/09-genealogia-sankey.png · 498 erros na base 1 (não é a "
              "base 3).")
    figura(s, "09-genealogia-sankey.png", 0.65, 1.95, 12.0, 4.9, "Sankey da genealogia da base 1",
           "analysis/2026-09-trace-law-flow/assets/")

    s = anexo("B", "As três bases lado a lado", "Fontes: base 1 = assets/08-overview-granularidade.png; bases 2 e 3 = "
              "figuras de overview da máquina da base (--figuras).")
    for i, (nome, rot, origem) in enumerate([
            ("08-overview-granularidade.png", "Base 1", "analysis/2026-09-trace-law-flow/assets/"),
            ("overview_base2.png", "Base 2", "--figuras; máquina da base 2"),
            ("overview_base3.png", "Base 3", "--figuras; máquina da base 3")]):
        text(s, rot, 0.65 + i * 4.05, 1.95, 3.9, 0.3, 13, TEAL, True)
        figura(s, nome, 0.65 + i * 4.05, 2.3, 3.9, 4.5, f"Overview · {rot}", origem)

    mapa = anexo("C", "Mapa das figuras: o que entrou e de onde veio", "Gerado por gerar_apresentacao.py — a mesma "
                 "lista está em figuras.md, ao lado do .pptx.", rodada=True)

    s = anexo("D", "O esquema das três bases", "Fontes: 2026-09-trace-law-flow/docs/05-schema.md; second/docs/"
              "02-relatorio-base2.md; third/docs/05-pedido-queries.md e 02-relatorio-achados.md §1.")
    tabela(s, [["", "Base 1", "Base 2", "Base 3"],
               ["Recorte", "out/25–ago/26 (11 meses)", "dez/25–ago/26 (9 meses; 88% jun–ago)", "nov/25–ago/26 (10 meses)"],
               ["Execuções", "1.000 (amostra)", "1.000 (amostra)", "71.881 (população)"],
               ["Passos (ActionSteps)", "5.781", "3.907", "283.925"],
               ["Erros estruturados", "498", "479", "30.141 (em 20.442 execuções)"],
               ["Unidade da linha", "um passo do trace", "um passo do trace", "um erro + a janela n−1 … n+2"],
               ["Prompt e modelo", "sim", "sim", "não"],
               ["Trajetória e resposta final", "sim", "sim", "não"],
               ["Denominador (todos os passos)", "sim (na amostra)", "sim (na amostra)", "sim (pedido de contagens)"],
               ["Falhas silenciosas", "medidas", "medidas", "fora (só erros com exceção)"]],
           0.65, 1.95, 12.0, 4.4, [3.0, 2.8, 3.2, 3.0], 11)

    s = anexo("E", "O sandbox, memória por memória (rodada de 21/10)",
              "Fontes: third/propostas/U_sandbox/dossie_base3.md §2h, itens 11 e 22; analysis/registros/memorias.json; "
              "analysis/decisoes.md 20–21/10.")
    tabela(s, [["Memória", "Erros", "Destino", "Degrau 2 (prompt da base 1, depois do Ajuste 28)"],
               ["C global (json, re, datetime)", "423", "memória global", "passa, menos no managerAgent (decisão a trazer)"],
               ["E global (depuração)", "317", "memória global", "passa (a repetição decidida pelo pesquisador)"],
               ["C do JoogleAnalytics (pd)", "231", "memória do papel", "passa"],
               ["B do OBFCivel (nomes do framework)", "44", "memória do papel", "falha conflito (alarme falso da regra) e "
                                                                            "repete — decisão a trazer"],
               ["D × 10 papéis (import fora da lista)", "40", "reforço, monitorada", "repetem a lista do prompt"],
               ["C do CalculoTrabalhista · B do RespostaBacen", "2 · 1", "raras, monitoradas", "falham recorrência"],
               ["F (como verificar no sandbox)", "37", "em aberto", "fora do catálogo"],
               ["sem nome", "5", "—", "—"]],
           0.65, 1.95, 12.0, 3.6, [3.7, 0.9, 2.1, 5.3], 10.5, direita=(1,))
    note(s, "Trajetória da unidade: 3.653 (14/10) → 1.436 (Ajuste 21, −2.217) → 1.401 (Ajuste 23, −35) → 1.100 "
            "(Ajustes 26 e 27, −167 −134).  “Uma causa só” é decidida pelo pesquisador, memória por memória.",
         y=5.75, h=0.55, size=10)

    s = anexo("F", "Uma linha da base 3: o erro e a janela [n−1] [n] [n+1] [n+2]",
              "Fonte: a fonte v2 da base 3, na máquina da base — anonimizado e varrido antes de sair (--figuras).")
    figura(s, "recorte_base3_janela.png", 0.65, 1.95, 12.0, 4.9, "Recorte real de uma linha da base 3",
           "--figuras; anonimizado, varredura de PII")

    s = anexo("G", "Um erro da base 1 em duas formas: a saída do drill_down e o trace",
              "Fonte: base 1 via drill_down.py — o caso do recorte do slide 2; anonimizado e varrido (--figuras).")
    figura(s, "recorte_base1_drilldown.png", 0.65, 1.95, 5.9, 4.9, "(a) A saída do drill_down.py",
           "--figuras; anonimizado, varredura de PII")
    figura(s, "recorte_base1_trace.png", 6.75, 1.95, 5.9, 4.9, "(b) O trace completo da execução",
           "--figuras; anonimizado, varredura de PII")

    s = anexo("H", "“Could not index”: um sintoma, três causas — o caso trabalhado",
              "Fonte: 09-metodologia-erro-a-memoria.md (o trio A/B/C); casos reais na máquina da base (--figuras).")
    figura(s, "trio_could_not_index.png", 0.65, 1.95, 12.0, 4.9, "Três erros reais com a mesma assinatura",
           "--figuras; a mensagem, a linha rejeitada, a saída do classify() e a do submecanismo(); anonimizado")

    s = anexo("I", "O pipeline etapa por etapa: os 9 passos da mineração",
              "Fontes: third/docs/06-procedimento-mineracao-candidatas.md; .claude/skills/mineracao-base/referencias/"
              "kit.md; o caso = dossiê do sandbox §2h, item 11.")
    tabela(s, [["Passo", "O que faz", "No sandbox (21/10)"],
               ["1 · rodada", "confere a rodada e o manifesto (hash de cada tabela)", "rodada 99e032ad"],
               ["2 · monitoramento", "o que já está em observação, antes de minerar", "—"],
               ["3 · parte genérica", "recorrência, sub-unidades, antes, depois, estabilidade", "1.100 erros · 982 "
                                                                                                "execuções · 10 meses"],
               ["4 · parte específica", "o script aprovado da unidade lê cada erro", "16 papéis · 6 listas autorizadas"],
               ["5 · confiança", "cobertura da regra (limiar 95%)", "99,6%"],
               ["6 · destino", "cada parte com o seu dono", "memória global, do papel, reforço, rara, em aberto"],
               ["7 · auditoria e encontro", "script independente recalcula; compara medida a medida", "77 de 77 iguais"],
               ["8 · arquivo final", "as memórias no esquema, validadas", "17 memórias no esquema 0.4"],
               ["9 · saída", "catálogo, manifesto, varredura de PII", "16 no catálogo + F fora · varredura limpa"]],
           0.65, 1.95, 12.0, 4.4, [2.6, 5.4, 4.0], 11)

    s = anexo("J", "A escada em detalhe: as checagens de cada degrau",
              "Fonte: 2026-09-trace-law-flow/docs/17-validacao-memorias.md; analysis/decisoes.md 21/10 (o degrau 2 "
              "pelo interpretador); pocs/blueprint-m1-pocs-minimos.md.")
    card(s, 0.65, 1.95, 5.9, 4.4, "DEGRAU 2 · CONFERIDA — 5 CHECAGENS POR SCRIPT",
         "**1 · Ambiente:** o que a memória manda existe? (o que o interpretador aceita — a mensagem de erro e os "
         "imports que rodaram —, não a lista do prompt)\n**2 · Conflito:** a memória contradiz o prompt?\n"
         "**3 · Repete:** a memória só repete o prompt?\n**4 · Recorrência:** ≥3 execuções e ≥2 meses\n"
         "**5 · Uma causa só:** decidida pelo pesquisador, memória por memória, com motivo", GOLD, 11, gap=6)
    card(s, 6.7, 1.95, 5.95, 2.1, "DEGRAU 3 · REVISADA — A LEITURA DE CASOS",
         "2 leitores às cegas leem casos sorteados por regra fixa; concordância medida; gate do pesquisador. Casos da "
         "base 1 onde houver; da base 2 onde não.", TEAL, 11)
    card(s, 6.7, 4.2, 5.95, 2.15, "DEGRAU 4 · VALIDADA — O EXPERIMENTO",
         "A POC mínima: reprodução do erro + compensação, antes × depois da memória, no mesmo recorte. Só aqui "
         "“candidata” vira “memória validada”.", RED, 11)

    s = anexo("K", "Destinos por unidade: qual unidade foi para qual destino",
              "Fonte: painel/ da rodada (executar.py painel).", rodada=True)
    figura(s, "unidades_destino_tipo.png", 0.65, 1.95, 12.0, 4.9, "Destinos por unidade",
           "painel/ da rodada (executar.py painel)")

    s = anexo("L", "Sinal de harness, o exemplo trabalhado: o prompt do ContestacaoCivel induz o erro",
              "Fonte: third/propostas/U_sandbox/dossie_base3.md §2 (parte A) e §2h item 4; base 2 e base 3.")
    card(s, 0.65, 1.95, 5.9, 2.6, "O QUE O PROMPT DIZ",
         "**Versão antiga (jun–ago):** “VERIFIQUE se a variável de saída já está preenchida em memória. Se já estiver, "
         "PULE a chamada” — manda verificar, não diz como.\n**Versão nova (só ago):** a mesma instrução + “não use "
         "globals()/locals(); use try/except NameError” — manda o como, e o como também quebra no sandbox.",
         GOLD, 10.5, gap=5)
    tabela(s, [["Base 2", "Verificam", "Com erro"],
               ["prompt antigo — globals()/locals()/vars", "36", "35"],
               ["prompt novo — try/except", "97", "97"]],
           6.7, 1.95, 5.95, 1.2, [3.85, 1.05, 1.05], 11, direita=(1, 2))
    tabela(s, [["Base 3 · ContestacaoCivel", "jul/26", "ago/26"],
               ["introspecção (globals, locals…)", "1.527", "709"],
               ["try/except", "95", "5.667"]],
           6.7, 3.3, 5.95, 1.2, [3.85, 1.05, 1.05], 11, direita=(1, 2))
    card(s, 0.65, 4.7, 12.0, 1.6, "A LEITURA",
         "Ensinou-se o jeito errado: o erro trocou de sintoma, não sumiu. Por isso o conserto é no prompt, não em "
         "memória — uma memória “não use globals” brigaria com a própria instrução.", RED, 13)

    s = anexo("M", "O prompt anuncia um ambiente que não é o real",
              "Fontes: third/propostas/U_sandbox/dossie_base3.md §2h, itens 16–21; analysis/decisoes.md 21/10; "
              "smolagents 1.26.0 (agents.py); pipeline/validacao/testar_sandbox_executor_proprio.py.")
    tabela(s, [["Papel (prompt sem json na lista)", "Base", "Passos com import json", "Rodou", "Barrado"],
               ["ContestacaoCivel", "2", "435", "435", "0"],
               ["RoteadorCivel", "2", "63", "63", "0"],
               ["OBFCivel", "2", "22", "22", "0"],
               ["3 papéis", "1", "11", "11", "0"],
               ["Total", "1 e 2", "531", "531", "0"]],
           0.65, 1.95, 6.6, 2.4, [2.7, 0.8, 1.4, 0.85, 0.85], 11, destaque=(5,), direita=(1, 2, 3, 4))
    tabela(s, [["Montagem no smolagents", "json no prompt", "json no interpretador", "import json"],
               ["padrão", "não", "não", "barrado"],
               ["executor= com json, agente sem", "não", "sim", "roda"],
               ["o inverso", "sim", "não", "barrado"]],
           7.4, 1.95, 5.25, 1.6, [2.15, 0.95, 1.15, 1.0], 10)
    card(s, 0.65, 4.55, 12.0, 1.75, "O QUE ISSO MUDA",
         "**No método:** o degrau 2 confere o ambiente pelo que o interpretador aceita, não pela lista do prompt.\n"
         "**Para a plataforma:** a montagem “executor= com json, agente sem” reproduz o que as bases mostram — é "
         "possível por configuração; se é assim na plataforma, só o código dela diz. Vai como pergunta no aviso.",
         TEAL, 11.5, gap=5)
    return mapa


def preencher_mapa(s):
    linhas = [["Slide", "Arquivo", "Estado", "De onde vem"]]
    for n, nome, caminho, origem in FIGURAS:
        linhas.append([f"{n:02}", nome, "inserida" if caminho else "FALTA", origem])
    tabela(s, linhas, 0.65, 1.95, 12.0, min(4.9, 0.25 * len(linhas)), [0.7, 3.2, 1.0, 7.1], 8.5)


# ---------------------------------------------------------------- montagem

def rodada_padrao():
    manifestos = sorted(RESULTADOS.glob("observada_*/manifesto.json"), key=lambda m: m.stat().st_mtime)
    return manifestos[-1].parent if manifestos else None


def montar(rodada, extra, saida):
    global R, CARIMBO
    pasta = saida.parent
    figs = pasta / "figuras"
    figs.mkdir(parents=True, exist_ok=True)
    BUSCA[:] = [p for p in [extra, figs] if p] + (
        [rodada / "relatorio", rodada / "painel", rodada / "taxas", rodada] if rodada else []) + [caminho_io(ASSETS_BASE1)]
    if rodada:
        m = json.loads((rodada / "manifesto.json").read_text(encoding="utf-8"))
        CARIMBO = f"rodada {m.get('rodada_id', rodada.name)[:16]} · {str(m.get('gerado_em', ''))[:10]}"
        erros = m.get("resumo", {}).get("erros_estruturados")
        if erros is not None and erros != ERROS_V2:
            AVISOS.append(f"A rodada tem {n_br(erros)} erros estruturados; o texto do deck é da fonte v2 "
                          f"({n_br(ERROS_V2)}). Rodada de ensaio?")
        if not any((rodada / sub).is_dir() for sub in ("relatorio", "painel", "taxas")):
            AVISOS.append(f"{rodada.name}/ sem relatorio/, painel/ e taxas/ — rode antes: executar.py relatorio "
                          f"--rodada {rodada}")
    R = Presentation()
    R.slide_width, R.slide_height = Inches(13.333), Inches(7.5)
    serie = serie_mensal(rodada)
    s01_abertura()
    s02_anatomia()
    s03_necessidade()
    s04_bases()
    s05_escala(serie, figs)
    s06_metodo()
    s07_destinos()
    s08_processo()
    s09_sandbox()
    s10_custo()
    s11_conclusoes()
    s12_pedido()
    s13_referencias()
    preencher_mapa(anexos())
    R.save(saida)
    faltam = [f for f in FIGURAS if f[2] is None]
    md = [f"# Figuras da apresentação — {CARIMBO or 'sem rodada (só marcadores e base 1)'}", "",
          f"`{saida.name}`: {len(R.slides)} slides; {len(FIGURAS) - len(faltam)} figuras inseridas, {len(faltam)} "
          "faltando.", "", "| Slide | Arquivo | Estado | De onde vem |", "|---|---|---|---|"]
    md += [f"| {n:02} | `{nome}` | {'inserida: ' + str(c) if c else '**FALTA**'} | {o} |" for n, nome, c, o in FIGURAS]
    if AVISOS:
        md += ["", "## Avisos", ""] + [f"- {a}" for a in AVISOS]
    (pasta / "figuras.md").write_text("\n".join(md) + "\n", encoding="utf-8")
    return len(R.slides), faltam


def main(argv=None):
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--rodada", type=Path, help="pasta resultados/observada_<id>/ (padrão: a mais recente)")
    p.add_argument("--figuras", type=Path, help="pasta com as figuras feitas à mão (nomes no anexo C)")
    p.add_argument("--saida", type=Path, help="o .pptx (padrão: <rodada>/apresentacao/apresentacao-stakeholders.pptx)")
    args = p.parse_args(argv)
    rodada = args.rodada or rodada_padrao()
    rodada = caminho_io(rodada) if rodada is not None else None
    if rodada is not None and not (rodada / "manifesto.json").is_file():
        raise SystemExit(f"{rodada}: sem manifesto.json — não é uma pasta de rodada.")
    saida = args.saida or ((rodada / "apresentacao") if rodada else (HERE / "apresentacao")) / \
        "apresentacao-stakeholders.pptx"
    saida = caminho_io(saida)
    saida.parent.mkdir(parents=True, exist_ok=True)
    total, faltam = montar(rodada, caminho_io(args.figuras) if args.figuras else None, saida)
    print(f"Apresentação: {saida} ({total} slides)")
    print(f"Figuras: {len(FIGURAS) - len(faltam)} inseridas, {len(faltam)} faltando — lista em {saida.parent / 'figuras.md'}")
    for nome in sorted({f[1] for f in faltam}):
        print(f"  FALTA {nome}")
    for a in AVISOS:
        print(f"AVISO: {a}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
